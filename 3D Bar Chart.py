
# import openpyxl module
import openpyxl

# import BarChart3D class from openpyxl.chart sub_module
from openpyxl.chart import BarChart3D,Reference

# Call a Workbook() function of openpyxl
# to create a new blank Workbook object
wb = openpyxl.Workbook()

# Get workbook active sheet
# from the active attribute.
sheet = wb.active

# write o to 9 in 1st column of the active sheet
for i in range(10):
    sheet.append([i])

values = Reference(sheet, min_col = 1, min_row = 1,
                         max_col = 1, max_row = 10)

# Create object of BarChart3D class
chart = BarChart3D()

chart.add_data(values)

# set the title of the chart
chart.title = " BAR-CHART3D "

# set the title of the x-axis
chart.x_axis.title = " X AXIS "

# set the title of the y-axis
chart.y_axis.title = " Y AXIS "

# add chart to the sheet
# the top-left corner of a chart
# is anchored to cell E2.
sheet.add_chart(chart, "E2")

# save the file
wb.save("BarChart3D.xlsx")
